from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from datetime import datetime
from typing import List, Dict
from pathlib import Path
from src.master_control.utils.logger import logger


class ExcelExporter:
    
    @staticmethod
    def export_records(records: List[Dict], output_path: str = None) -> str:
        if output_path is None:
            timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
            output_path = f"test_records_{timestamp}.xlsx"
        
        wb = Workbook()
        ws = wb.active
        ws.title = "测试记录"
        
        headers = [
            "ID", "SN", "IP", "测试时间", "总体状态",
            "POE状态", "POE信息", "POE详情",
            "硬盘状态", "硬盘信息", "硬盘详情",
            "WiFi状态", "WiFi信息", "WiFi详情",
            "蓝牙状态", "蓝牙信息", "蓝牙详情",
            "星闪状态", "星闪信息", "星闪详情",
            "4G状态", "4G信息", "4G详情"
        ]
        
        header_fill = PatternFill(start_color="366092", end_color="366092", fill_type="solid")
        header_font = Font(bold=True, color="FFFFFF", size=11)
        header_alignment = Alignment(horizontal="center", vertical="center")
        
        border = Border(
            left=Side(style='thin'),
            right=Side(style='thin'),
            top=Side(style='thin'),
            bottom=Side(style='thin')
        )
        
        for col_num, header in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col_num)
            cell.value = header
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = header_alignment
            cell.border = border
        
        for row_num, record in enumerate(records, 2):
            ws.cell(row=row_num, column=1, value=record.get('id', ''))
            ws.cell(row=row_num, column=2, value=record.get('sn', ''))
            ws.cell(row=row_num, column=3, value=record.get('ip', ''))
            ws.cell(row=row_num, column=4, value=record.get('test_time', ''))
            ws.cell(row=row_num, column=5, value=record.get('overall_status', ''))
            
            ws.cell(row=row_num, column=6, value=record.get('poe_status', ''))
            ws.cell(row=row_num, column=7, value=record.get('poe_message', ''))
            ws.cell(row=row_num, column=8, value=record.get('poe_detail', ''))
            
            ws.cell(row=row_num, column=9, value=record.get('disk_status', ''))
            ws.cell(row=row_num, column=10, value=record.get('disk_message', ''))
            ws.cell(row=row_num, column=11, value=record.get('disk_detail', ''))
            
            ws.cell(row=row_num, column=12, value=record.get('wifi_status', ''))
            ws.cell(row=row_num, column=13, value=record.get('wifi_message', ''))
            ws.cell(row=row_num, column=14, value=record.get('wifi_detail', ''))
            
            ws.cell(row=row_num, column=15, value=record.get('bluetooth_status', ''))
            ws.cell(row=row_num, column=16, value=record.get('bluetooth_message', ''))
            ws.cell(row=row_num, column=17, value=record.get('bluetooth_detail', ''))
            
            ws.cell(row=row_num, column=18, value=record.get('nearlink_status', ''))
            ws.cell(row=row_num, column=19, value=record.get('nearlink_message', ''))
            ws.cell(row=row_num, column=20, value=record.get('nearlink_detail', ''))
            
            ws.cell(row=row_num, column=21, value=record.get('lte_4g_status', ''))
            ws.cell(row=row_num, column=22, value=record.get('lte_4g_message', ''))
            ws.cell(row=row_num, column=23, value=record.get('lte_4g_detail', ''))
            
            for col_num in range(1, len(headers) + 1):
                cell = ws.cell(row=row_num, column=col_num)
                cell.border = border
                cell.alignment = Alignment(vertical="center")
        
        for col_num in range(1, len(headers) + 1):
            ws.column_dimensions[ws.cell(row=1, column=col_num).column_letter].width = 15
        
        ws.column_dimensions['B'].width = 20
        ws.column_dimensions['C'].width = 18
        ws.column_dimensions['D'].width = 20
        
        wb.save(output_path)
        logger.info(f"Excel 导出成功: {output_path}")
        
        return output_path
